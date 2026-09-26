"""Спецификация двигателя из каталога модели: PDF на русском и английском и книга Excel.

Запуск из папки engine:
    node tools/dump-catalog.mjs > /tmp/catalog.json
    python3 tools/build_spec.py /tmp/catalog.json spec [папка со шрифтами IBM Plex]

Шрифты: PlexSans-Regular.ttf, PlexSans-SemiBold.ttf, PlexSans-Bold.ttf, PlexMono-Regular.ttf (IBM Plex)
и RobotoCond-600.ttf (Roboto Condensed) с Google Fonts. Без них используется DejaVu Sans.
"""
import json
import os
import sys
from datetime import date
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Frame, KeepTogether, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)

INK = colors.HexColor('#11151a')
INK2 = colors.HexColor('#4b5563')
INK3 = colors.HexColor('#7a8491')
LINE = colors.HexColor('#d9dde3')
ACCENT = colors.HexColor('#0b66d6')
TI = colors.HexColor('#8a6a2a')
ROW = colors.HexColor('#f4f6f8')

TXT = {
    'ru': {
        'title': 'Спецификация деталей', 'eyebrow': 'Интерактивная 3D-модель · спецификация',
        'engine': 'Двигатель', 'cols': ['Поз.', 'Деталь', 'Кол-во', 'Материал', 'Назначение', 'Параметры'],
        'asm': 'Узел', 'parts': 'поз.', 'page': 'Стр.', 'summary': 'Состав',
        'summaryCols': ['Узел', 'Название', 'Позиций'], 'total': 'Всего',
        'date': 'Сформировано', 'sheet': 'Спецификация', 'how': 'Как работает',
        'asmCol': 'Узел', 'nameCol': 'Деталь',
    },
    'en': {
        'title': 'Parts specification', 'eyebrow': 'Interactive 3D model · specification',
        'engine': 'Engine', 'cols': ['Item', 'Part', 'Qty', 'Material', 'Purpose', 'Specifications'],
        'asm': 'Assembly', 'parts': 'items', 'page': 'Page', 'summary': 'Contents',
        'summaryCols': ['No.', 'Assembly', 'Items'], 'total': 'Total',
        'date': 'Generated', 'sheet': 'Specification', 'how': 'How it works',
        'asmCol': 'Assembly', 'nameCol': 'Part',
    },
}


def fonts(fdir):
    names = {
        'Sans': 'PlexSans-Regular.ttf', 'SansSB': 'PlexSans-SemiBold.ttf',
        'Cond': 'RobotoCond-600.ttf', 'CondB': 'PlexSans-Bold.ttf', 'Mono': 'PlexMono-Regular.ttf',
    }
    dv = '/usr/share/fonts/truetype/dejavu/'
    fallback = {'Sans': 'DejaVuSans.ttf', 'SansSB': 'DejaVuSans-Bold.ttf', 'Cond': 'DejaVuSans-Bold.ttf',
                'CondB': 'DejaVuSans-Bold.ttf', 'Mono': 'DejaVuSansMono.ttf'}
    for k, f in names.items():
        p = os.path.join(fdir or '', f)
        pdfmetrics.registerFont(TTFont(k, p if fdir and os.path.exists(p) else dv + fallback[k]))


def rich(s):
    """Текст для Paragraph: экранирование и верхний индекс вместо символов, которых нет в шрифте."""
    return escape(s).replace('⁻¹', '<super>−1</super>').replace('‑', '-')


def pdf(cat, lang, out):
    T, E = TXT[lang], cat['engine'][lang]
    st = {
        'eyebrow': ParagraphStyle('eb', fontName='Cond', fontSize=8.5, leading=11, textColor=INK3, spaceAfter=2),
        'h1': ParagraphStyle('h1', fontName='CondB', fontSize=30, leading=32, textColor=INK),
        'kind': ParagraphStyle('k', fontName='Sans', fontSize=12, leading=16, textColor=INK2, spaceBefore=4),
        'lead': ParagraphStyle('l', fontName='Sans', fontSize=10.5, leading=15, textColor=INK, spaceBefore=8),
        'h2': ParagraphStyle('h2', fontName='CondB', fontSize=15, leading=18, textColor=INK),
        'about': ParagraphStyle('a', fontName='Sans', fontSize=9, leading=12.5, textColor=INK2, spaceBefore=3, spaceAfter=6),
        'cell': ParagraphStyle('c', fontName='Sans', fontSize=8, leading=10.4, textColor=INK, alignment=TA_LEFT),
        'cellB': ParagraphStyle('cb', fontName='SansSB', fontSize=8.4, leading=10.6, textColor=INK),
        'cellM': ParagraphStyle('cm', fontName='Mono', fontSize=7.6, leading=10, textColor=TI),
        'cellS': ParagraphStyle('cs', fontName='Sans', fontSize=7.4, leading=9.6, textColor=INK2),
        'head': ParagraphStyle('hd', fontName='Cond', fontSize=7.8, leading=10, textColor=INK3),
        'note': ParagraphStyle('n', fontName='Sans', fontSize=8, leading=11, textColor=INK3, spaceBefore=10),
        'spec': ParagraphStyle('sp', fontName='Sans', fontSize=9, leading=12, textColor=INK2),
        'specV': ParagraphStyle('sv', fontName='Mono', fontSize=8.6, leading=12, textColor=INK),
    }
    W, H = landscape(A4)
    M = 12 * mm

    def deco(c, doc):
        c.saveState()
        c.setStrokeColor(LINE)
        c.setLineWidth(0.6)
        c.line(M, H - M + 3 * mm, W - M, H - M + 3 * mm)
        c.setFont('Cond', 8)
        c.setFillColor(INK3)
        c.drawString(M, H - M + 5 * mm, f'V8 5.0 · {T["title"]}')
        c.drawRightString(W - M, M - 6 * mm, f'{T["page"]} {doc.page}')
        c.drawString(M, M - 6 * mm, f'{T["date"]}: {date.today().isoformat()}')
        c.restoreState()

    doc = BaseDocTemplate(out, pagesize=(W, H), leftMargin=M, rightMargin=M, topMargin=M + 4 * mm, bottomMargin=M,
                          title=f'V8 5.0 — {T["title"]}', author='V8 5.0', subject=T['title'])
    frame = Frame(M, M, W - 2 * M, H - 2 * M - 4 * mm, id='f', leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id='p', frames=[frame], onPage=deco)])

    story = []
    # ---- титульный разворот: описание, характеристики, состав
    left = [Paragraph(escape(T['eyebrow']).upper(), st['eyebrow']), Paragraph(escape(E['name']), st['h1']),
            Paragraph(escape(E['kind']), st['kind']), Paragraph(rich(E['intro']), st['lead']), Spacer(1, 8)]
    spec = Table([[Paragraph(rich(k), st['spec']), Paragraph(rich(v), st['specV'])] for k, v in E['specs']],
                 colWidths=[48 * mm, 80 * mm])
    spec.setStyle(TableStyle([('LINEBELOW', (0, 0), (-1, -1), 0.5, LINE), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                              ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                              ('LEFTPADDING', (0, 0), (-1, -1), 0), ('LINEABOVE', (0, 0), (-1, 0), 0.5, LINE)]))
    left += [spec, Paragraph(rich(E['note']), st['note'])]
    total = sum(len(a['parts']) for a in cat['asms'])
    rows = [[Paragraph(escape(h).upper(), st['head']) for h in T['summaryCols']]]
    for a in cat['asms']:
        rows.append([Paragraph(a['no'], st['cellM']), Paragraph(escape(a[lang]['name']), st['cellB']), Paragraph(str(len(a['parts'])), st['cell'])])
    rows.append([Paragraph('', st['cell']), Paragraph(escape(T['total']), st['cellB']), Paragraph(str(total), st['cellB'])])
    summ = Table(rows, colWidths=[12 * mm, 88 * mm, 16 * mm])
    summ.setStyle(TableStyle([('LINEBELOW', (0, 0), (-1, -1), 0.5, LINE), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                              ('TOPPADDING', (0, 0), (-1, -1), 3.2), ('BOTTOMPADDING', (0, 0), (-1, -1), 3.2),
                              ('LEFTPADDING', (0, 0), (-1, -1), 0), ('LINEBELOW', (0, 0), (-1, 0), 0.8, INK3),
                              ('LINEABOVE', (0, -1), (-1, -1), 0.8, INK3)]))
    right = [Paragraph(escape(T['summary']).upper(), st['eyebrow']), Spacer(1, 4), summ]
    cover = Table([[left, right]], colWidths=[140 * mm, W - 2 * M - 140 * mm])
    cover.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0),
                               ('RIGHTPADDING', (0, 0), (0, 0), 14 * mm), ('RIGHTPADDING', (1, 0), (1, 0), 0)]))
    story.append(cover)

    # ---- узлы
    cw = [15 * mm, 44 * mm, 24 * mm, 46 * mm, 80 * mm]
    cw.append(W - 2 * M - sum(cw))
    for a in cat['asms']:
        story.append(CondPageBreak(60 * mm))
        story.append(Spacer(1, 10))
        hdr = [Paragraph(f'{escape(T["asm"]).upper()} {a["no"]} · {len(a["parts"])} {escape(T["parts"])}', st['eyebrow']),
               Paragraph(escape(a[lang]['name']), st['h2']), Paragraph(rich(a[lang]['about']), st['about'])]
        rows = [[Paragraph(escape(h).upper(), st['head']) for h in T['cols']]]
        for p in a['parts']:
            d = p[lang]
            params = '<br/>'.join(f'{rich(k)}: <font name="Mono" color="#11151a">{rich(v)}</font>' for k, v in d.get('data') or [])
            rows.append([Paragraph(p['no'], st['cellM']), Paragraph(rich(d['name']), st['cellB']), Paragraph(rich(d['qty']), st['cell']),
                         Paragraph(rich(d['mat']), st['cell']), Paragraph(rich(d['role']), st['cell']), Paragraph(params, st['cellS'])])
        t = Table(rows, colWidths=cw, repeatRows=1)
        style = [('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LINEBELOW', (0, 0), (-1, -1), 0.5, LINE),
                 ('LINEBELOW', (0, 0), (-1, 0), 0.8, INK3), ('TOPPADDING', (0, 0), (-1, -1), 4),
                 ('BOTTOMPADDING', (0, 0), (-1, -1), 4), ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 6)]
        for i in range(2, len(rows), 2):
            style.append(('BACKGROUND', (0, i), (-1, i), ROW))
        t.setStyle(TableStyle(style))
        story.append(KeepTogether(hdr + [t]) if len(rows) <= 6 else KeepTogether(hdr))
        if len(rows) > 6:
            story.append(t)
    doc.build(story)


def xlsx(cat, out):
    wb = Workbook()
    thin = Side(style='thin', color='D9DDE3')
    head_fill = PatternFill('solid', fgColor='11151A')
    grp_fill = PatternFill('solid', fgColor='E8EEF7')
    wrap = Alignment(wrap_text=True, vertical='top')

    # лист с характеристиками на двух языках
    ws = wb.active
    ws.title = 'Двигатель · Engine'
    ws['A1'] = 'V8 5.0'
    ws['A1'].font = Font(size=20, bold=True)
    ws['A2'] = f'{cat["engine"]["ru"]["kind"]} / {cat["engine"]["en"]["kind"]}'
    ws['A2'].font = Font(size=11, color='4B5563')
    ws.append([])
    ws.append(['Параметр', 'Значение', 'Parameter', 'Value'])
    for c in ws[4]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head_fill
    for (k, v), (ke, ve) in zip(cat['engine']['ru']['specs'], cat['engine']['en']['specs']):
        ws.append([k, v, ke, ve])
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = wrap
            c.border = Border(bottom=thin)
    ws.append([])
    ws.append([cat['engine']['ru']['note']])
    ws.append([cat['engine']['en']['note']])
    for col, w in zip('ABCD', (26, 44, 26, 44)):
        ws.column_dimensions[col].width = w

    for lang in ('ru', 'en'):
        T = TXT[lang]
        ws = wb.create_sheet(T['sheet'])
        cols = [T['cols'][0], T['asmCol'], T['nameCol'], T['cols'][2], T['cols'][3], T['cols'][4], T['how'], T['cols'][5]]
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = head_fill
            c.alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 22
        for a in cat['asms']:
            ws.append([a['no'], a[lang]['name']])
            r = ws.max_row
            for c in ws[r]:
                c.font = Font(bold=True)
            for ci in range(1, len(cols) + 1):
                ws.cell(r, ci).fill = grp_fill
            for p in a['parts']:
                d = p[lang]
                params = '\n'.join(f'{k}: {v}' for k, v in d.get('data') or [])
                ws.append([p['no'], a[lang]['name'], d['name'], d['qty'], d['mat'], d['role'], d['how'], params])
                for c in ws[ws.max_row]:
                    c.alignment = wrap
                    c.border = Border(bottom=thin)
        for i, w in enumerate((8, 26, 30, 14, 30, 48, 70, 40), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{ws.max_row}'
    wb.properties.title = 'V8 5.0 — спецификация / specification'
    wb.save(out)


if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    fonts(sys.argv[3] if len(sys.argv) > 3 else None)
    cat = json.load(open(src, encoding='utf-8'))
    os.makedirs(dst, exist_ok=True)
    for lang in ('ru', 'en'):
        pdf(cat, lang, os.path.join(dst, f'v8-spec-{lang}.pdf'))
    xlsx(cat, os.path.join(dst, 'v8-spec.xlsx'))
    print('ok')
